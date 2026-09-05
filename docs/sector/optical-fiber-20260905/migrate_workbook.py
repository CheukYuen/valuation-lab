"""One-time archive migration. Normal HTML builds never require the workbook."""
import json, hashlib, datetime
from pathlib import Path
from zipfile import ZipFile
import openpyxl
B=Path(__file__).resolve().parent
p=B.parents[2]/'outputs/sector-overview-20260905/行业数据与估值附录.xlsx'
a=openpyxl.load_workbook(p,data_only=False);b=openpyxl.load_workbook(p,data_only=True)
def value(v): return v.isoformat() if isinstance(v,(datetime.datetime,datetime.date)) else v
sheets=[]
for s in a:
 cells=[]
 for row in s:
  for c in row:
   if c.value is None and not c.comment:continue
   cells.append(dict(ref=c.coordinate,value=value(b[s.title][c.coordinate].value),formula=c.value[1:] if c.data_type=='f' else None,comment=c.comment.text if c.comment else None,format=c.number_format))
 sheets.append(dict(name=s.title,cells=cells,merges=[str(x) for x in s.merged_cells.ranges]))

# Presentation is extracted alongside the values: freeze panes, widths, heights and a
# de-duplicated palette of fill / font / alignment, indexed per cell.
def rgb(c):
 if c is None or getattr(c,'type',None)!='rgb':return None
 v=c.rgb
 if not isinstance(v,str) or len(v)!=8 or v[:2]=='00':return None
 return v[2:].upper()
def style(c):
 f=c.fill
 fill=rgb(f.fgColor) if f is not None and f.patternType=='solid' else None
 fo=c.font;al=c.alignment
 return dict(fill=None if fill=='FFFFFF' else fill,color=rgb(fo.color),bold=bool(fo.b),italic=bool(fo.i),
             size=float(fo.sz) if fo.sz else None,name=fo.name,
             halign=al.horizontal,valign=al.vertical,wrap=bool(al.wrap_text))
DEFAULT=dict(fill=None,color=None,bold=False,italic=False,size=None,name=None,halign=None,valign=None,wrap=False)
palette=[DEFAULT];seen={json.dumps(DEFAULT,sort_keys=True):0};layout=[]
for ws in a:
 styles={}
 for r in range(1,ws.max_row+1):
  for col in range(1,ws.max_column+1):
   c=ws.cell(row=r,column=col);st=style(c)
   if st==DEFAULT:continue
   k=json.dumps(st,sort_keys=True,ensure_ascii=False)
   if k not in seen:seen[k]=len(palette);palette.append(st)
   styles[c.coordinate]=seen[k]
 layout.append(dict(
  name=ws.title,max_row=ws.max_row,max_col=ws.max_column,
  freeze=ws.freeze_panes,grid_lines=bool(ws.sheet_view.showGridLines),
  default_row_height=ws.sheet_format.defaultRowHeight,default_col_width=ws.sheet_format.defaultColWidth,
  col_width={k:round(v.width,2) for k,v in sorted(ws.column_dimensions.items()) if v.width},
  row_height={str(k):v.height for k,v in sorted(ws.row_dimensions.items()) if v.height},
  hidden_cols=[k for k,v in ws.column_dimensions.items() if v.hidden],
  hidden_rows=[str(k) for k,v in ws.row_dimensions.items() if v.hidden],
  styles=styles))

with ZipFile(p) as z: charts={n:z.read(n).decode() for n in z.namelist() if '/charts/chart' in n and n.endswith('.xml')}
out=dict(archive_sha256=hashlib.sha256(p.read_bytes()).hexdigest(),sheets=sheets,original_charts=charts,
         style=dict(palette=palette,sheets=layout))
out['counts']={k:sum(test(c) for s in sheets for c in s['cells']) for k,test in [('cells',lambda c:True),('formulas',lambda c:bool(c['formula'])),('comments',lambda c:bool(c['comment']))]}
out['counts']['palette']=len(palette)
out['counts']['grid_cells']=sum(s['max_row']*s['max_col'] for s in layout)
(B/'workbook_data.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n')
print(out['counts'])
